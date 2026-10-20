"""엑셀 양식·가져오기·내보내기 테스트 (가상 데이터)."""
import io
from datetime import date

import pytest
from openpyxl import load_workbook

from pmo.demo import seed_demo
from pmo.excel_io import SHEETS, build_import, export_bytes, template_bytes
from pmo.service import Service, ValidationError

TODAY = date(2026, 10, 19)


@pytest.fixture
def service(tmp_path):
    return Service(tmp_path / "import.db", clock=lambda: TODAY)


def workbook(**sheets) -> bytes:
    """양식에 시트별 행(dict: 열 이름 → 값)을 채운 파일."""
    wb = load_workbook(io.BytesIO(template_bytes()))
    for sheet, rows in sheets.items():
        ws = wb[sheet]
        names = [name for name, _, _ in SHEETS[sheet]]
        for r, row in enumerate(rows, 2):
            for name, value in row.items():
                ws.cell(row=r, column=names.index(name) + 1, value=value)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


GOOD = dict(
    인력=[{"인력키": "P01", "이름": "가상일", "업체명": "가상업체", "분류": "협력"},
        {"인력키": "P02", "이름": "가상이", "업체명": "가상업체", "분류": "정", "계정 ID": "fake2"}],
    투입=[{"인력키": "P01", "투입 과제": "과제A", "시작일": date(2026, 4, 27), "종료일": "2026-05-08", "등급": "L6",
         "투입 Type": "Full time", "투입 MM": 9, "구매시스템 MM": "0.445"},
        {"인력키": "P02", "시작일": "2026-11-01", "종료일": "2026-11-30", "등급": "L5", "투입 Type": "Part time",
         "투입 MM": 0.5, "매출단가": 2000000, "단가 기준 단위": "원/MM"}],
    구매계약=[{"인력키": "P01", "업체명": "가상업체", "계약번호": "C-1", "계약 시작일": "2026-04-27",
           "계약 종료일": "2026-10-31", "상태": "확정"},
          {"인력키": "P01", "업체명": "가상업체", "계약 시작일": "2026-11-01", "계약 종료일": "2027-03-31",
           "상태": "준비", "이전 계약번호": "C-1", "연결 투입 과제": "과제A"}],
)


def test_template_has_sheets_required_marks_and_dropdowns():
    wb = load_workbook(io.BytesIO(template_bytes()))
    assert wb.sheetnames == ["안내", "인력", "투입", "구매계약"]
    assert [c.value for c in wb["인력"][1]] == ["인력키 *", "이름 *", "업체명 *", "분류 *", "계정 ID"]
    assert any("정,협력,컨설팅" in v.formula1 for v in wb["인력"].data_validations.dataValidation)
    assert wb["투입"].max_row == 1                                   # 빈 양식(예시 데이터 없음)


def test_import_saves_everything_in_one_go(service):
    plan = build_import(service, workbook(**GOOD))
    assert not plan["errors"], plan["errors"]
    assert plan["summary"] == {"인력": 2, "투입": 2, "구매계약": 2}
    assert plan["mm_check"][0]["일치"] == "예"                         # 구매시스템 MM 0.445와 같음
    result = service.import_batch(plan, token="imp-1", actor="담당자")
    assert result == {"인력": 2, "투입": 2, "구매계약": 2}
    p1 = next(p for p in service.people() if p["name"] == "가상일")
    full = service.assignments(p1["id"])[0]
    assert full["mm"] == "0.445"                                     # Full time은 입력 MM(9)을 무시하고 계산
    assert len(service.checklist(full["id"])) == 4                   # 근무 환경 준비 항목 생성
    contracts = sorted(service.contracts(p1["id"]), key=lambda c: c["start_date"])
    assert contracts[1]["previous_contract_id"] == contracts[0]["id"]
    assert contracts[1]["assignment_id"] == full["id"]
    assert service.import_batch(plan, token="imp-1", actor="담당자") == result   # 재실행해도 중복 없음
    assert len(service.people()) == 2
    log = [r for r in service.run_log() if r["operation"] == "import_batch"][0]
    assert "가상일" not in log["input_json"]                           # 실행 기록에는 건수만


def test_errors_block_import_and_nothing_is_saved(service):
    bad = dict(
        인력=[{"인력키": "P01", "이름": "가상일", "업체명": "가상업체", "분류": "협력"},
            {"인력키": "P01", "이름": "중복키", "업체명": "가상업체", "분류": "협력"},
            {"인력키": "P03", "이름": "가상삼", "업체명": "가상업체", "분류": "직원"},
            {"인력키": "P04", "이름": "가상사", "업체명": "가상업체", "분류": "정"}],
        투입=[{"인력키": "P01", "시작일": "2026-11-30", "종료일": "2026-11-01", "등급": "L6", "투입 Type": "Full time"},
            {"인력키": "X9", "시작일": "2026-11-01", "종료일": "2026-11-30", "등급": "L6", "투입 Type": "Full time"},
            {"인력키": "P01", "시작일": "2026-11-01", "종료일": "2026-11-30", "등급": "L9", "투입 Type": "Full time"}],
        구매계약=[{"인력키": "P04", "업체명": "가상업체", "계약 시작일": "2026-11-01", "계약 종료일": "2026-12-31", "상태": "준비"},
              {"인력키": "P01", "업체명": "가상업체", "계약 시작일": "2026-11-01", "계약 종료일": "2026-12-31", "상태": "확정"}],
    )
    plan = build_import(service, workbook(**bad))
    messages = " / ".join(f"{e['시트']}{e['행']}:{e['내용']}" for e in plan["errors"])
    for expected in ("인력3:인력키 'P01'가 파일 안에서 중복", "인력4:분류는", "투입2:종료일은", "투입3:인력키 'X9'",
                     "투입4:허용된 등급", "구매계약2:구매계약은 분류가", "구매계약3:확정 계약은 계약번호"):
        assert expected in messages, messages
    with pytest.raises(ValidationError, match="오류가 있는 파일"):
        service.import_batch(plan, token="imp-bad", actor="담당자")
    assert not service.people()


def test_purchase_mm_mismatch_and_overlap_are_reported(service):
    rows = dict(인력=GOOD["인력"][:1],
               투입=[{"인력키": "P01", "시작일": "2026-04-27", "종료일": "2026-05-08", "등급": "L6",
                    "투입 Type": "Full time", "구매시스템 MM": 0.404},
                   {"인력키": "P01", "시작일": "2026-05-01", "종료일": "2026-05-31", "등급": "L6",
                    "투입 Type": "Full time"}])
    plan = build_import(service, workbook(**rows))
    assert plan["mm_check"] == [{"행": 2, "인력키": "P01", "기간": "2026-04-27~2026-05-08",
                                 "앱 계산": "0.445", "구매시스템": "0.404", "일치": "아니오"}]
    assert any("겹칩니다" in w["내용"] for w in plan["warnings"])
    assert not plan["errors"]


def test_existing_people_by_id_and_namesake_guard(service):
    seed_demo(service, TODAY)
    rows = dict(인력=[{"인력키": "N1", "이름": "이가상", "업체명": "베타컨설팅", "분류": "컨설팅"}],
               투입=[{"인력키": "DEMO-002", "투입 과제": "새 과제", "시작일": "2027-01-04", "종료일": "2027-01-29",
                    "등급": "L6", "투입 Type": "Full time"}])
    plan = build_import(service, workbook(**rows))
    assert any("이미 등록된 인력과 이름·업체가 같습니다(DEMO-002)" in e["내용"] for e in plan["errors"])
    plan = build_import(service, workbook(투입=rows["투입"]))
    assert not plan["errors"]
    service.import_batch(plan, token="imp-existing", actor="담당자")
    assert any(a["project"] == "새 과제" for a in service.assignments("DEMO-002"))


def test_export_round_trip_adds_nothing(service):
    seed_demo(service, TODAY)
    exported = export_bytes(service)
    wb = load_workbook(io.BytesIO(exported))
    assert wb["인력"].max_row - 1 == len(service.people())
    assert wb["구매계약"].max_row - 1 == len(service.contracts())
    plan = build_import(service, exported)
    assert not plan["errors"], plan["errors"][:3]
    assert plan["summary"] == {"인력": 0, "투입": 0, "구매계약": 0}
    assert len(plan["skipped"]) == len(service.people()) + len(service.assignments()) + len(service.contracts())


def test_not_an_excel_file(service):
    with pytest.raises(ValidationError, match="엑셀 파일"):
        build_import(service, b"not excel")


def test_command_line_import_with_backup(tmp_path, monkeypatch, capsys):
    import sys
    from scripts import import_excel
    db = tmp_path / "pmo.db"
    file = tmp_path / "in.xlsx"
    file.write_bytes(workbook(**GOOD))
    monkeypatch.setenv("PMO_DB_PATH", str(db))
    Service(db)                                                       # 빈 업무 DB
    monkeypatch.setattr(sys, "argv", ["import_excel", str(file), "--yes", "--export", str(tmp_path / "out.xlsx")])
    assert import_excel.main() == 0
    out = capsys.readouterr().out
    assert "저장했습니다: 인력 2명, 투입 2건, 구매계약 2건" in out
    assert len(list((tmp_path / "backups").glob("pmo-before-import-*.db"))) == 1
    assert (tmp_path / "out.xlsx").exists()
    assert len(Service(db).people()) == 2
    assert import_excel.main() == 0                                   # 다시 실행해도 중복 없음
    assert "이미 가져온 파일입니다(인력 2명" in capsys.readouterr().out


def test_command_line_stops_on_errors(tmp_path, monkeypatch, capsys):
    import sys
    from scripts import import_excel
    db = tmp_path / "pmo.db"
    file = tmp_path / "bad.xlsx"
    file.write_bytes(workbook(인력=[{"인력키": "P01", "이름": "가상", "분류": "협력"}]))
    monkeypatch.setenv("PMO_DB_PATH", str(db))
    monkeypatch.setattr(sys, "argv", ["import_excel", str(file), "--yes"])
    assert import_excel.main() == 1
    assert "오류가 있어 저장하지 않았습니다" in capsys.readouterr().out
    assert not Service(db).people()
