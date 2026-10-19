"""엑셀 양식·가져오기·내보내기 (PC 안에서만 처리, 외부 전송 없음).

- template_bytes(): 빈 양식(안내·인력·투입·구매계약 시트)
- build_import(service, data): 파일을 읽어 검증하고 가져오기 계획을 만든다. DB는 바꾸지 않는다.
  오류가 하나라도 있으면 가져올 수 없다. 이미 등록된 투입·계약과 같은 줄은 건너뛴다(내보낸 파일을 다시 넣어도 중복 없음).
- 저장은 Service.import_batch(plan, ...)가 한 트랜잭션으로 한다(전부 저장되거나 전혀 저장되지 않음).
- export_bytes(service): 현재 DB를 같은 양식으로 내려받는다. 인력키 = 내부 관리번호.
"""
import hashlib
import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from pmo.config import grades
from pmo.contracts import CONTRACT_CATEGORIES, CONTRACT_STATUSES, validate_contract
from pmo.errors import ValidationError
from pmo.service import Service, validate_assignment, validate_person

CATEGORIES = ("정", "협력", "컨설팅")
TYPES = ("Full time", "Part time")

# (열 이름, 필수, 데이터 키) — 시트 순서·열 순서가 곧 양식이다.
SHEETS = {
    "인력": [("인력키", True, "key"), ("이름", True, "name"), ("업체명", True, "company"),
           ("분류", True, "category"), ("계정 ID", False, "account_id")],
    "투입": [("인력키", True, "key"), ("투입 과제", False, "project"), ("매출귀속 프로젝트", False, "revenue_project"),
           ("시작일", True, "start_date"), ("종료일", True, "end_date"), ("등급", True, "grade"),
           ("투입 Type", True, "assignment_type"), ("투입 MM", False, "mm"),
           ("구매시스템 MM", False, "purchase_mm"), ("매출단가", False, "sales_rate"),
           ("매입단가", False, "purchase_rate"), ("단가 기준 단위", False, "rate_unit")],
    "구매계약": [("인력키", True, "key"), ("업체명", True, "company"), ("계약번호", False, "contract_no"),
             ("계약 시작일", True, "start_date"), ("계약 종료일", True, "end_date"), ("상태", True, "status"),
             ("연결 투입 과제", False, "assignment_project"), ("이전 계약번호", False, "previous_no")],
}
DATE_KEYS = {"start_date", "end_date"}
GUIDE = [
    "CoE 인력관리 데이터 가져오기 양식",
    "",
    "1. '인력', '투입', '구매계약' 시트의 2행부터 한 줄에 한 건씩 입력합니다. 열 이름은 바꾸지 마세요.",
    "2. 열 이름 뒤의 * 는 필수입니다. 모르는 값은 비워 두면 '미입력'으로 저장됩니다.",
    "3. 인력키: 파일 안에서 사람을 구분하는 값입니다(예: P01). 투입·구매계약 시트에서 같은 인력키로 연결합니다.",
    "   이미 앱에 등록된 사람에게 투입·계약을 더하려면 인력키에 그 사람의 내부 관리번호를 넣고 '인력' 시트에는 적지 않습니다.",
    "4. 날짜는 YYYY-MM-DD 형식 또는 엑셀 날짜로 입력합니다.",
    "5. 투입 MM: Part time만 입력합니다. Full time은 앱이 구매시스템 방식으로 계산합니다.",
    "   구매시스템 MM: 비교용(선택). 넣으면 앱 계산값과 다른 줄을 가져오기 전에 알려줍니다.",
    "6. 단가를 넣으면 '단가 기준 단위'(예: 원/MM)도 넣습니다.",
    "7. 구매계약은 분류가 협력·컨설팅인 사람만 등록합니다. 상태가 '확정'이면 계약번호가 필요합니다.",
    "   연결 투입 과제: 그 사람의 투입이 여러 건이면 어느 투입의 계약인지 투입 과제명으로 적습니다.",
    "   이전 계약번호: 연장 계약이면 이전 계약의 계약번호를 적습니다.",
    "8. 앱의 '데이터 가져오기·내보내기' 화면에서 파일을 올리면 미리보기로 오류와 건수를 먼저 보여주고, 확인해야 저장합니다.",
    "   이미 등록된 투입·계약과 같은 줄은 건너뜁니다.",
    "",
    "실제 인력·계약 데이터가 든 파일은 PC 안에서만 다루고 Git이나 외부로 보내지 않습니다.",
]

_HEAD_FILL = PatternFill("solid", fgColor="0F1B33")
_REQ_FILL = PatternFill("solid", fgColor="1D4ED8")


def _header(name: str, required: bool) -> str:
    return f"{name} *" if required else name


def _new_workbook() -> Workbook:
    wb = Workbook()
    guide = wb.active
    guide.title = "안내"
    for i, line in enumerate(GUIDE, 1):
        guide.cell(row=i, column=1, value=line)
    guide["A1"].font = Font(bold=True, size=14)
    guide.column_dimensions["A"].width = 120
    for sheet, columns in SHEETS.items():
        ws = wb.create_sheet(sheet)
        for col, (name, required, key) in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col, value=_header(name, required))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = _REQ_FILL if required else _HEAD_FILL
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[cell.column_letter].width = max(12, len(name) * 2 + 4)
        ws.freeze_panes = "A2"
        lists = {"category": CATEGORIES, "grade": grades(), "assignment_type": TYPES,
                 "status": CONTRACT_STATUSES}
        for col, (_, _, key) in enumerate(columns, 1):
            if key in lists:
                letter = ws.cell(row=1, column=col).column_letter
                rule = DataValidation(type="list", formula1='"' + ",".join(lists[key]) + '"', allow_blank=True)
                rule.add(f"{letter}2:{letter}1000")
                ws.add_data_validation(rule)
    return wb


def _to_bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def template_bytes() -> bytes:
    return _to_bytes(_new_workbook())


# ---------------------------------------------------------------------------
# 읽기·검증
# ---------------------------------------------------------------------------

def _cell(value, key: str) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        value = format(Decimal(repr(value)).normalize(), "f")
    text = str(value).strip()
    if key in DATE_KEYS and len(text) >= 10 and text[4] in "-./":
        text = text[:10].replace(".", "-").replace("/", "-")
    return text or None


def read_workbook(data: bytes) -> tuple[dict[str, list[dict]], list[dict]]:
    """시트별 [{"row": 엑셀 행 번호, 데이터 키: 값}]과 구조 오류."""
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - 어떤 이유든 엑셀로 열 수 없으면 사용자에게 알린다
        raise ValidationError(f"엑셀 파일(.xlsx)을 열 수 없습니다: {type(exc).__name__}") from exc
    sheets, errors = {}, []
    for sheet, columns in SHEETS.items():
        sheets[sheet] = []
        if sheet not in wb.sheetnames:
            errors.append({"시트": sheet, "행": "-", "내용": "시트가 없습니다. 양식을 내려받아 사용하세요."})
            continue
        rows = wb[sheet].iter_rows(values_only=True)
        header = [str(h or "").replace("*", "").strip() for h in next(rows, ())]
        index = {}
        for name, required, key in columns:
            if name in header:
                index[key] = header.index(name)
            elif required:
                errors.append({"시트": sheet, "행": 1, "내용": f"'{name}' 열이 없습니다."})
        for number, values in enumerate(rows, 2):
            item = {key: _cell(values[i] if i < len(values) else None, key) for key, i in index.items()}
            if any(v is not None for v in item.values()):
                sheets[sheet].append(dict(item, row=number))
    wb.close()
    return sheets, errors


def _overlap(a: dict, b: dict) -> bool:
    return a["start_date"] <= b["end_date"] and b["start_date"] <= a["end_date"]


def build_import(service: Service, data: bytes) -> dict:
    """파일을 검증해 가져오기 계획을 만든다(저장 안 함)."""
    sheets, errors = read_workbook(data)
    warnings, skipped, mm_check = [], [], []

    def err(sheet, row, text):
        errors.append({"시트": sheet, "행": row, "내용": text})

    db_people = {p["id"]: p for p in service.people()}
    db_assignments = service.assignments()
    db_contracts = service.contracts()

    # --- 인력 ---
    people, categories = [], {pid: p["category"] for pid, p in db_people.items()}
    for row in sheets.get("인력", []):
        key = row.get("key")
        if key in db_people:
            skipped.append({"시트": "인력", "행": row["row"], "내용": f"이미 등록된 인력({key})이라 건너뜁니다."})
            continue
        try:
            if not key:
                raise ValidationError("인력키를 입력하세요.")
            if key in categories:
                raise ValidationError(f"인력키 '{key}'가 파일 안에서 중복됩니다.")
            person = validate_person({k: row.get(k) for k in ("name", "company", "category", "account_id")})
            if not person["company"]:
                raise ValidationError("업체명을 입력하세요.")
            same = [p for p in db_people.values() if p["name"] == person["name"] and p["company"] == person["company"]]
            if same:
                raise ValidationError(f"이미 등록된 인력과 이름·업체가 같습니다({same[0]['id']}). 같은 사람이면 "
                                      "'인력' 시트에서 빼고 투입·계약 시트의 인력키에 그 관리번호를 넣으세요.")
        except ValidationError as exc:
            err("인력", row["row"], str(exc))
            continue
        categories[key] = person["category"]
        people.append({"row": row["row"], "key": key, "data": person})
    new_keys = {p["key"] for p in people}

    def known(sheet, row) -> bool:
        key = row.get("key")
        if key in new_keys or key in db_people:
            return True
        err(sheet, row["row"], f"인력키 '{key or ''}'를 '인력' 시트나 등록된 관리번호에서 찾을 수 없습니다.")
        return False

    # --- 투입 ---
    assignments = []
    for row in sheets.get("투입", []):
        if not known("투입", row):
            continue
        try:
            fields = {k: row.get(k) for k in ("project", "revenue_project", "start_date", "end_date", "grade",
                                              "assignment_type", "mm", "sales_rate", "purchase_rate", "rate_unit")}
            if fields["assignment_type"] == "Full time":
                fields["mm"] = None                      # Full time은 앱이 계산한다
            value = validate_assignment(dict(fields, person_id=row["key"]))
            value.pop("person_id")
        except ValidationError as exc:
            err("투입", row["row"], str(exc))
            continue
        same = [a for a in db_assignments if a["person_id"] == row["key"] and a["start_date"] == value["start_date"]
                and a["end_date"] == value["end_date"] and (a["project"] or None) == value["project"]]
        if same:
            skipped.append({"시트": "투입", "행": row["row"], "내용": "이미 등록된 투입이라 건너뜁니다."})
            continue
        others = [a for a in db_assignments if a["person_id"] == row["key"]] + \
                 [a["data"] for a in assignments if a["key"] == row["key"]]
        if any(_overlap(value, o) for o in others):
            warnings.append({"시트": "투입", "행": row["row"], "내용": "같은 사람의 다른 투입과 기간이 겹칩니다(검토 대상)."})
        purchase = row.get("purchase_mm")
        if purchase is not None and value["assignment_type"] == "Full time":
            try:
                same_mm = Decimal(purchase) == Decimal(value["mm"])
            except InvalidOperation:
                err("투입", row["row"], "구매시스템 MM은 숫자로 입력하세요.")
                continue
            mm_check.append({"행": row["row"], "인력키": row["key"], "기간": f"{value['start_date']}~{value['end_date']}",
                             "앱 계산": value["mm"], "구매시스템": purchase, "일치": "예" if same_mm else "아니오"})
        assignments.append({"row": row["row"], "key": row["key"], "data": value})

    # --- 구매계약 ---
    contracts = []
    for row in sheets.get("구매계약", []):
        if not known("구매계약", row):
            continue
        key = row["key"]
        try:
            if categories.get(key) not in CONTRACT_CATEGORIES:
                raise ValidationError("구매계약은 분류가 협력·컨설팅인 인력만 등록합니다.")
            value = validate_contract({k: row.get(k) for k in ("company", "contract_no", "start_date", "end_date",
                                                                "status")} | {"person_id": key})
            value.pop("person_id")
            assignment_ref = _find_assignment(key, row.get("assignment_project"), assignments, db_assignments)
            previous_ref = _find_previous(key, row.get("previous_no"), contracts, db_contracts)
        except ValidationError as exc:
            err("구매계약", row["row"], str(exc))
            continue
        same = [c for c in db_contracts if c["person_id"] == key and c["start_date"] == value["start_date"]
                and c["end_date"] == value["end_date"] and (c["contract_no"] or None) == value["contract_no"]]
        if same:
            skipped.append({"시트": "구매계약", "행": row["row"], "내용": "이미 등록된 계약이라 건너뜁니다."})
            continue
        contracts.append({"row": row["row"], "key": key, "data": value,
                          "assignment": assignment_ref, "previous": previous_ref})

    return {
        "file_hash": hashlib.sha256(data).hexdigest(),
        "people": people, "assignments": assignments, "contracts": contracts,
        "errors": errors, "warnings": warnings, "skipped": skipped, "mm_check": mm_check,
        "summary": {"인력": len(people), "투입": len(assignments), "구매계약": len(contracts)},
    }


def _find_assignment(key, project, new_rows, db_rows):
    """연결 투입: 과제명으로 찾는다. 비어 있으면 그 사람의 투입이 하나일 때만 자동 연결."""
    new = [i for i, a in enumerate(new_rows) if a["key"] == key]
    db = [a for a in db_rows if a["person_id"] == key]
    if project:
        new = [i for i in new if new_rows[i]["data"]["project"] == project]
        db = [a for a in db if a["project"] == project]
        if len(new) + len(db) != 1:
            raise ValidationError(f"연결 투입 과제 '{project}'에 맞는 투입이 {len(new) + len(db)}건입니다. 한 건이어야 합니다.")
    elif len(new) + len(db) != 1:
        return None                                  # 연결하지 않음(선택 항목)
    return ["new", new[0]] if new else ["db", db[0]["id"]]


def _find_previous(key, number, new_rows, db_rows):
    if not number:
        return None
    new = [i for i, c in enumerate(new_rows) if c["key"] == key and c["data"]["contract_no"] == number]
    db = [c for c in db_rows if c["person_id"] == key and c["contract_no"] == number]
    if len(new) + len(db) != 1:
        raise ValidationError(f"이전 계약번호 '{number}'를 같은 사람의 계약에서 한 건으로 찾을 수 없습니다. "
                              "이전 계약을 이 줄보다 위에 적었는지 확인하세요.")
    return ["new", new[0]] if new else ["db", db[0]["id"]]


# ---------------------------------------------------------------------------
# 내보내기
# ---------------------------------------------------------------------------

def export_bytes(service: Service) -> bytes:
    """현재 DB를 가져오기 양식과 같은 모양으로 내려받는다(인력키 = 내부 관리번호)."""
    wb = _new_workbook()
    wb["안내"].cell(row=len(GUIDE) + 2, column=1,
                  value=f"내보낸 시각: {datetime.now().isoformat(timespec='seconds')} · DB: {service.path.name}")
    assignments = service.assignments()
    by_assignment = {a["id"]: a for a in assignments}
    contracts = service.contracts()
    by_contract = {c["id"]: c for c in contracts}
    rows = {
        "인력": [{"key": p["id"], **p} for p in service.people()],
        "투입": [{"key": a["person_id"], **a} for a in assignments],
        "구매계약": [{"key": c["person_id"], **c,
                   "assignment_project": (by_assignment.get(c.get("assignment_id")) or {}).get("project"),
                   "previous_no": (by_contract.get(c.get("previous_contract_id")) or {}).get("contract_no")}
                  for c in contracts],
    }
    for sheet, columns in SHEETS.items():
        ws = wb[sheet]
        for r, item in enumerate(rows[sheet], 2):
            for c, (_, _, key) in enumerate(columns, 1):
                value = item.get(key)
                cell = ws.cell(row=r, column=c, value=date.fromisoformat(value) if key in DATE_KEYS and value else value)
                if key in DATE_KEYS and value:
                    cell.number_format = "yyyy-mm-dd"
    return _to_bytes(wb)
